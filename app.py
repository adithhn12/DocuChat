import hashlib
import json
import os
import sqlite3
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import docx
import pandas as pd
from pptx import Presentation
from PyPDF2 import PdfReader

import streamlit as st

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:
    from langchain.text_splitter import RecursiveCharacterTextSplitter

from langchain_community.embeddings import HuggingFaceEmbeddings, OllamaEmbeddings
from langchain_community.llms import Ollama
from langchain_community.vectorstores import FAISS

# ------------------------------------------------------------------------------
# 1. Database Operations (User Authentication)
# ------------------------------------------------------------------------------
def init_db():
    conn = sqlite3.connect('users.db')
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users
                 (username TEXT PRIMARY KEY, password TEXT)''')
    conn.commit()
    conn.close()

init_db()

def register_user(username, password):
    conn = sqlite3.connect('users.db')
    c = conn.cursor()
    try:
        c.execute("INSERT INTO users VALUES (?, ?)", (username, password))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()

def authenticate_user(username, password):
    conn = sqlite3.connect('users.db')
    c = conn.cursor()
    c.execute("SELECT * FROM users WHERE username=? AND password=?", (username, password))
    result = c.fetchone()
    conn.close()
    return result is not None

# ------------------------------------------------------------------------------
# 2. Session Management
# ------------------------------------------------------------------------------
def init_session():
    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False
    if "username" not in st.session_state:
        st.session_state.username = None
    if "chat_history" not in st.session_state:
        st.session_state.chat_history = []
    if "vectorstore" not in st.session_state:
        st.session_state.vectorstore = None
    if "processed_docs" not in st.session_state:
        st.session_state.processed_docs = []

# ------------------------------------------------------------------------------
# 3. Document Extraction & Parsing
# ------------------------------------------------------------------------------
def extract_pdf_text(pdf_file):
    try:
        return " ".join(page.extract_text() or "" for page in PdfReader(pdf_file).pages)
    except Exception as e:
        st.error(f"Error reading PDF {pdf_file.name}: {str(e)}")
        return ""

def extract_docx_text(docx_file):
    try:
        doc = docx.Document(docx_file)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        st.error(f"Error reading DOCX {docx_file.name}: {str(e)}")
        return ""

def extract_pptx_text(pptx_file):
    try:
        prs = Presentation(pptx_file)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)
    except Exception as e:
        st.error(f"Error reading PPTX {pptx_file.name}: {str(e)}")
        return ""

def extract_xlsx_text(xlsx_file):
    try:
        df = pd.read_excel(xlsx_file, sheet_name=None)
        text = []
        for sheet_name, sheet_data in df.items():
            text.append(f"Sheet: {sheet_name}")
            text.append(sheet_data.to_string())
        return "\n".join(text)
    except Exception as e:
        st.error(f"Error reading XLSX {xlsx_file.name}: {str(e)}")
        return ""

def extract_txt_text(txt_file):
    try:
        return txt_file.read().decode("utf-8")
    except Exception as e:
        st.error(f"Error reading TXT {txt_file.name}: {str(e)}")
        return ""

def extract_text(file):
    file_type = file.name.split(".")[-1].lower()
    if file_type == "pdf":
        return extract_pdf_text(file)
    elif file_type == "docx":
        return extract_docx_text(file)
    elif file_type == "pptx":
        return extract_pptx_text(file)
    elif file_type == "xlsx":
        return extract_xlsx_text(file)
    elif file_type == "txt":
        return extract_txt_text(file)
    else:
        st.error(f"Unsupported file type: {file_type}")
        return ""

def process_files_parallel(files):
    with ThreadPoolExecutor(max_workers=4) as executor:
        return " ".join(executor.map(extract_text, files))

def chunk_text(text):
    return RecursiveCharacterTextSplitter(
        chunk_size=768,
        chunk_overlap=100,
        length_function=len
    ).split_text(text)

# ------------------------------------------------------------------------------
# 4. Groq API Custom Direct Client
# ------------------------------------------------------------------------------
class GroqLLM:
    def __init__(self, api_key, model_name="llama-3.1-8b-instant", temperature=0.2):
        self.api_key = api_key
        self.model_name = model_name
        self.temperature = temperature

    def invoke(self, prompt):
        url = "https://api.groq.com/openai/v1/chat/completions"
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json",
            "User-Agent": "DocuChat/1.0"
        }
        payload = {
            "model": self.model_name,
            "messages": [
                {"role": "user", "content": prompt}
            ],
            "temperature": self.temperature
        }
        req = urllib.request.Request(
            url, 
            data=json.dumps(payload).encode("utf-8"), 
            headers=headers, 
            method="POST"
        )
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                res_json = json.loads(resp.read().decode("utf-8"))
                return res_json["choices"][0]["message"]["content"]
        except urllib.error.HTTPError as e:
            err_body = e.read().decode("utf-8")
            try:
                err_json = json.loads(err_body)
                msg = err_json.get("error", {}).get("message", err_body)
            except Exception:
                msg = err_body
            raise RuntimeError(f"Groq API Error ({e.code}): {msg}")
        except Exception as e:
            raise RuntimeError(f"Connection Error: {str(e)}")

# ------------------------------------------------------------------------------
# 5. Embeddings, Vector Stores & LLM
# ------------------------------------------------------------------------------
def is_ollama_available(host_url="http://localhost:11434"):
    try:
        tags_url = f"{host_url.rstrip('/')}/api/tags"
        with urllib.request.urlopen(tags_url, timeout=1.5) as resp:
            return resp.status == 200
    except Exception:
        return False

def get_embeddings():
    embed_model = st.session_state.get("ollama_embed_model", "nomic-embed-text")
    host_url = st.session_state.get("ollama_host", "http://localhost:11434")
    if is_ollama_available(host_url) and embed_model != "huggingface":
        try:
            return OllamaEmbeddings(base_url=host_url, model=embed_model)
        except Exception:
            pass
    return HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

def get_llm():
    provider = st.session_state.get("ai_provider", "Groq Cloud API (Recommended for Cloud)")
    
    # 1. Groq Cloud API Engine
    if provider.startswith("Groq"):
        groq_key = st.session_state.get("groq_api_key", "").strip()
        if not groq_key and "GROQ_API_KEY" in st.secrets:
            groq_key = st.secrets["GROQ_API_KEY"]
        if not groq_key and os.environ.get("GROQ_API_KEY"):
            groq_key = os.environ.get("GROQ_API_KEY")
            
        if groq_key:
            model_name = st.session_state.get("groq_model", "llama-3.1-8b-instant")
            return GroqLLM(api_key=groq_key, model_name=model_name, temperature=0.2)
        else:
            return None
            
    # 2. Ollama Local / Custom Host Engine
    else:
        host_url = st.session_state.get("ollama_host", "http://localhost:11434")
        model_name = st.session_state.get("ollama_model", "llama3.2")
        if is_ollama_available(host_url):
            try:
                return Ollama(base_url=host_url, model=model_name)
            except Exception:
                return None
        else:
            return None

def get_vectorstore(text_chunks):
    embed_model_name = st.session_state.get("ollama_embed_model", "nomic-embed-text")
    cache_key = hashlib.md5(f"{embed_model_name}_{''.join(text_chunks)}".encode()).hexdigest()
    cache_path = f"vector_cache/{st.session_state.username}_{cache_key}"
    embeddings = get_embeddings()
    
    if os.path.exists(cache_path):
        try:
            vs = FAISS.load_local(
                cache_path, 
                embeddings,
                allow_dangerous_deserialization=True
            )
            test_emb = embeddings.embed_query("test")
            if vs.index.d == len(test_emb):
                return vs
            else:
                import shutil
                shutil.rmtree(cache_path, ignore_errors=True)
        except Exception:
            if os.path.exists(cache_path):
                import shutil
                shutil.rmtree(cache_path, ignore_errors=True)
    
    vectorstore = FAISS.from_texts(text_chunks, embedding=embeddings)
    os.makedirs("vector_cache", exist_ok=True)
    vectorstore.save_local(cache_path)
    return vectorstore

def format_human_prose(query, docs):
    import re
    paragraphs = []
    for doc in docs:
        txt = doc.page_content.strip()
        if txt:
            cleaned = re.sub(r'\s+', ' ', txt)
            paragraphs.append(cleaned)
    combined = " ".join(paragraphs)
    
    if len(combined) > 700:
        combined = combined[:700] + "..."
        
    return f"Based on your document, here is what I found:\n\n{combined}"

# ------------------------------------------------------------------------------
# 6. UI Custom Styling (CSS Injection)
# ------------------------------------------------------------------------------
def inject_custom_css():
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@500;600;700;800&family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Plus Jakarta Sans', sans-serif;
    }

    .stApp {
        background: radial-gradient(circle at 15% 15%, rgba(99, 102, 241, 0.12), transparent 45%),
                    radial-gradient(circle at 85% 85%, rgba(168, 85, 247, 0.12), transparent 45%),
                    #090d16 !important;
        color: #f1f5f9;
    }

    /* Sidebar Glassmorphism */
    [data-testid="stSidebar"] {
        background: rgba(15, 23, 42, 0.75) !important;
        backdrop-filter: blur(16px);
        border-right: 1px solid rgba(255, 255, 255, 0.08);
    }
    
    [data-testid="stSidebar"] h1, [data-testid="stSidebar"] h2, [data-testid="stSidebar"] h3 {
        font-family: 'Outfit', sans-serif;
        color: #f8fafc;
    }

    /* Input text styling */
    .stTextInput input, .stSelectbox select {
        background-color: rgba(30, 41, 59, 0.8) !important;
        border: 1px solid rgba(255, 255, 255, 0.12) !important;
        color: #f8fafc !important;
        border-radius: 10px !important;
        padding: 10px 14px !important;
        font-size: 0.95rem !important;
    }
    
    .stTextInput input:focus, .stSelectbox select:focus {
        border-color: #8b5cf6 !important;
        box-shadow: 0 0 12px rgba(139, 92, 246, 0.3) !important;
    }

    /* Modern Streamlit Buttons */
    .stButton > button {
        background: linear-gradient(135deg, #6366f1 0%, #8b5cf6 100%) !important;
        color: #ffffff !important;
        border: none !important;
        border-radius: 12px !important;
        padding: 10px 20px !important;
        font-weight: 600 !important;
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        transition: all 0.3s ease !important;
        box-shadow: 0 4px 15px rgba(99, 102, 241, 0.3) !important;
        width: 100%;
    }

    .stButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(139, 92, 246, 0.5) !important;
        background: linear-gradient(135deg, #4f46e5 0%, #7c3aed 100%) !important;
    }

    /* Gradient Hero Title */
    .hero-title {
        font-family: 'Outfit', sans-serif;
        font-size: 2.5rem;
        font-weight: 800;
        background: linear-gradient(135deg, #818cf8 0%, #c084fc 50%, #f472b6 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0.2rem;
    }
    
    .hero-subtitle {
        color: #94a3b8;
        font-size: 1.02rem;
        margin-bottom: 1.2rem;
    }

    /* Status Badge Pill */
    .status-badge {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        padding: 6px 14px;
        border-radius: 20px;
        background: rgba(99, 102, 241, 0.15);
        border: 1px solid rgba(99, 102, 241, 0.3);
        color: #a5b4fc;
        font-size: 0.85rem;
        font-weight: 600;
        margin-right: 8px;
        margin-bottom: 10px;
    }

    /* Doc Pill Tags */
    .doc-pill {
        display: inline-flex;
        align-items: center;
        gap: 6px;
        padding: 6px 14px;
        border-radius: 10px;
        background: rgba(30, 41, 59, 0.8);
        border: 1px solid rgba(139, 92, 246, 0.3);
        color: #e2e8f0;
        font-size: 0.88rem;
        margin: 4px 6px 4px 0;
        font-weight: 500;
    }

    /* Chat Messages Layout */
    .chat-bubble-container {
        display: flex;
        gap: 12px;
        margin: 16px 0;
        align-items: flex-start;
    }

    .chat-bubble-container.user {
        flex-direction: row-reverse;
    }

    .avatar-icon {
        width: 40px;
        height: 40px;
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 1.2rem;
        flex-shrink: 0;
    }

    .avatar-icon.user {
        background: linear-gradient(135deg, #6366f1, #a855f7);
        color: white;
        box-shadow: 0 0 12px rgba(168, 85, 247, 0.4);
    }

    .avatar-icon.bot {
        background: rgba(30, 41, 59, 0.9);
        border: 1px solid #8b5cf6;
        color: #c084fc;
        box-shadow: 0 0 12px rgba(139, 92, 246, 0.3);
    }

    .chat-bubble {
        max-width: 78%;
        padding: 14px 18px;
        border-radius: 16px;
        line-height: 1.65;
        font-size: 0.96rem;
    }

    .chat-bubble.user {
        background: linear-gradient(135deg, #4338ca 0%, #6d28d9 100%);
        color: #ffffff;
        border-bottom-right-radius: 4px;
        box-shadow: 0 4px 16px rgba(79, 70, 229, 0.35);
    }

    .chat-bubble.bot {
        background: rgba(30, 41, 59, 0.85);
        color: #f1f5f9;
        border: 1px solid rgba(255, 255, 255, 0.08);
        border-left: 4px solid #8b5cf6;
        border-bottom-left-radius: 4px;
        backdrop-filter: blur(10px);
        box-shadow: 0 4px 18px rgba(0, 0, 0, 0.25);
    }

    .chat-timestamp {
        font-size: 0.72rem;
        margin-top: 8px;
        opacity: 0.7;
        text-align: right;
    }

    /* Auth Card Box */
    .auth-box {
        background: rgba(30, 41, 59, 0.7);
        backdrop-filter: blur(16px);
        border: 1px solid rgba(255, 255, 255, 0.1);
        border-radius: 20px;
        padding: 36px;
        box-shadow: 0 15px 35px rgba(0, 0, 0, 0.4);
    }

    /* Chat Input Bar */
    [data-testid="stChatInput"] {
        border-radius: 16px !important;
        border: 1px solid rgba(139, 92, 246, 0.4) !important;
        background: rgba(30, 41, 59, 0.9) !important;
        backdrop-filter: blur(12px) !important;
        box-shadow: 0 4px 24px rgba(0, 0, 0, 0.35) !important;
    }
    </style>
    """, unsafe_allow_html=True)

# ------------------------------------------------------------------------------
# 7. Auth Views (Login & Registration)
# ------------------------------------------------------------------------------
def login_page():
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("""
        <div style="text-align: center; margin-top: 40px; margin-bottom: 24px;">
            <div class="hero-title">DocuChat AI</div>
            <div class="hero-subtitle">Private Local & Cloud Document Intelligence</div>
        </div>
        """, unsafe_allow_html=True)
        
        username = st.text_input("Username", key="login_user")
        password = st.text_input("Password", type="password", key="login_pass")
        st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)
        
        if st.button("🔐 Sign In", use_container_width=True):
            if authenticate_user(username, password):
                st.session_state.authenticated = True
                st.session_state.username = username
                st.rerun()
            else:
                st.error("Invalid username or password")
        
        st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
        if st.button("✨ Create New Account", use_container_width=True):
            st.session_state.show_register = True
            st.rerun()

def register_page():
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown("""
        <div style="text-align: center; margin-top: 40px; margin-bottom: 24px;">
            <div class="hero-title">Create Account</div>
            <div class="hero-subtitle">Get started with DocuChat AI</div>
        </div>
        """, unsafe_allow_html=True)
        
        username = st.text_input("Username", key="reg_user")
        password = st.text_input("Password", type="password", key="reg_pass")
        confirm_password = st.text_input("Confirm Password", type="password", key="reg_confirm")
        st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)
        
        if st.button("🚀 Register Account", use_container_width=True):
            if password != confirm_password:
                st.error("Passwords don't match")
            elif len(username) < 3:
                st.error("Username must be at least 3 characters")
            elif len(password) < 6:
                st.error("Password must be at least 6 characters")
            elif register_user(username, password):
                st.success("Registration successful! Please login.")
                st.session_state.show_register = False
                st.rerun()
            else:
                st.error("Username already exists")
        
        st.markdown("<div style='height: 10px;'></div>", unsafe_allow_html=True)
        if st.button("⬅️ Back to Login", use_container_width=True):
            st.session_state.show_register = False
            st.rerun()

# ------------------------------------------------------------------------------
# 8. Query Handling & Generation
# ------------------------------------------------------------------------------
def handle_query(query):
    if not st.session_state.get("vectorstore"):
        st.warning("Please process documents first")
        return
    
    with st.spinner("Analyzing documents..."):
        try:
            docs = st.session_state.vectorstore.similarity_search(query, k=3)
            if not docs:
                answer = "I searched your uploaded documents, but couldn't find any relevant information for your query."
            else:
                context_str = "\n".join([doc.page_content.strip() for doc in docs if doc.page_content.strip()])
                prompt = f"Answer the user query in a natural, friendly, conversational way based strictly on the context below.\n\nContext:\n{context_str}\n\nQuestion: {query}\n\nAnswer:"

                llm = get_llm()
                if llm is not None:
                    try:
                        res = llm.invoke(prompt)
                        if hasattr(res, "content"):
                            answer = res.content.strip()
                        elif isinstance(res, str):
                            answer = res.strip()
                        else:
                            answer = str(res).strip()
                    except Exception as err:
                        answer = f"⚠️ **Groq AI Generation Error**: {str(err)}\n\n{format_human_prose(query, docs)}"
                else:
                    provider = st.session_state.get("ai_provider", "Groq")
                    if provider.startswith("Groq"):
                        answer = f"⚠️ **Missing Groq API Key**: Please enter your Groq API Key in the sidebar settings.\n\n{format_human_prose(query, docs)}"
                    else:
                        answer = format_human_prose(query, docs)
                
            st.session_state.chat_history.append((query, answer))
            st.rerun()
        except Exception as e:
            err_msg = str(e) or repr(e)
            if "AssertionError" in err_msg or "dimension" in err_msg.lower():
                st.error("Vector index dimension mismatch. Please click 'Process Documents' in the sidebar to re-index your documents.")
            else:
                st.error(f"Error generating response: {err_msg}")

# ------------------------------------------------------------------------------
# 9. Main Chat Application Interface
# ------------------------------------------------------------------------------
def chat_interface():
    with st.sidebar:
        st.title("⚙️ AI Engine Settings")
        
        st.session_state.ai_provider = st.sidebar.selectbox(
            "AI Provider",
            ["Groq Cloud API (Recommended for Cloud)", "Ollama (Local / Custom Host)"],
            index=0 if "GROQ_API_KEY" in st.secrets or st.session_state.get("groq_api_key") else 0
        )
        
        if st.session_state.ai_provider.startswith("Groq"):
            st.session_state.groq_api_key = st.sidebar.text_input(
                "Groq API Key",
                type="password",
                value=st.session_state.get("groq_api_key", ""),
                help="Get your free API key at console.groq.com"
            )
            st.session_state.groq_model = st.sidebar.selectbox(
                "Groq Model Name",
                ["llama-3.1-8b-instant", "llama3-8b-8192", "llama3-70b-8192", "mixtral-8x7b-32768", "gemma2-9b-it"],
                index=0
            )
        else:
            st.session_state.ollama_host = st.sidebar.text_input(
                "Ollama Host URL",
                value=st.session_state.get("ollama_host", "http://localhost:11434"),
                help="Local or remote Ollama server endpoint (e.g. http://localhost:11434 or ngrok URL)"
            )
            st.session_state.ollama_model = st.sidebar.text_input(
                "Ollama Model Name",
                value=st.session_state.get("ollama_model", "llama3.2"),
                help="Name of pulled Ollama model (e.g. llama3.2, mistral, llama2, qwen2.5)"
            )
        
        st.session_state.ollama_embed_model = st.sidebar.text_input(
            "Embedding Model",
            value=st.session_state.get("ollama_embed_model", "nomic-embed-text"),
            help="Name of Ollama embedding model or enter 'huggingface' for local HF embeddings"
        )

        st.markdown("---")
        st.subheader("📄 Document Processing")
        uploaded_files = st.file_uploader(
            "Upload Documents", 
            accept_multiple_files=True,
            type=["pdf", "docx", "pptx", "xlsx", "txt"]
        )
        
        if st.button("🚀 Process Documents"):
            if uploaded_files:
                with st.status("Processing documents...", expanded=True) as status:
                    st.write("Extracting text...")
                    text = process_files_parallel(uploaded_files)
                    
                    if not text.strip():
                        st.error("No text could be extracted")
                        return
                    
                    st.write("Creating chunks...")
                    chunks = chunk_text(text)
                    
                    st.write("Building vector store...")
                    st.session_state.vectorstore = get_vectorstore(chunks)
                    st.session_state.processed_docs = [file.name for file in uploaded_files]
                    
                    st.write("Ready for questions!")
                    st.success("Documents processed successfully!")
            else:
                st.warning("Please upload documents first")
        
        st.markdown("<div style='height: 20px;'></div>", unsafe_allow_html=True)
        if st.button("🚪 Logout"):
            st.session_state.authenticated = False
            st.session_state.username = None
            st.rerun()
    
    if st.session_state.get("ai_provider", "").startswith("Groq"):
        active_model = f"{st.session_state.get('groq_model', 'llama-3.1-8b-instant')} (Groq Cloud API)"
    else:
        active_model = f"{st.session_state.get('ollama_model', 'llama3.2')} (Ollama)"
        
    embed_model = st.session_state.get('ollama_embed_model', 'nomic-embed-text')
    
    st.markdown(f"""
    <div style="margin-bottom: 24px;">
        <div class="hero-title">DocuChat AI</div>
        <div class="hero-subtitle">Ask questions & gain insights from your documents locally or via Cloud AI</div>
        <div>
            <span class="status-badge">⚡ Model: {active_model}</span>
            <span class="status-badge">🧠 Embeddings: {embed_model}</span>
            <span class="status-badge">👤 User: {st.session_state.username}</span>
        </div>
    </div>
    """, unsafe_allow_html=True)
    
    if st.session_state.processed_docs:
        pills_html = "".join([f'<span class="doc-pill">📄 {doc}</span>' for doc in st.session_state.processed_docs])
        st.markdown(f"""
        <div style="margin-bottom: 20px;">
            <div style="font-weight: 600; color: #94a3b8; font-size: 0.82rem; letter-spacing: 0.05em; margin-bottom: 6px;">PROCESSED DOCUMENTS</div>
            <div>{pills_html}</div>
        </div>
        """, unsafe_allow_html=True)
    
    if st.session_state.vectorstore:
        st.markdown("<div style='font-weight: 600; color: #94a3b8; font-size: 0.82rem; letter-spacing: 0.05em; margin-bottom: 8px;'>SUGGESTED QUESTIONS</div>", unsafe_allow_html=True)
        cols = st.columns(3)
        questions = [
            "What is this document about?",
            "Summarize the key points",
            "What are the main conclusions?"
        ]
        for i, q in enumerate(questions):
            if cols[i%3].button(q):
                handle_query(q)

    chat_container = st.container()
    with chat_container:
        for query_text, answer_text in st.session_state.chat_history:
            time_str = datetime.now().strftime('%H:%M')
            # User Bubble
            st.markdown(f"""
            <div class="chat-bubble-container user">
                <div class="avatar-icon user">👤</div>
                <div class="chat-bubble user">
                    <div>{query_text}</div>
                    <div class="chat-timestamp">{time_str}</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
            # Bot Bubble
            st.markdown(f"""
            <div class="chat-bubble-container bot">
                <div class="avatar-icon bot">🤖</div>
                <div class="chat-bubble bot">
                    <div>{answer_text}</div>
                    <div class="chat-timestamp">{time_str} • {active_model}</div>
                </div>
            </div>
            """, unsafe_allow_html=True)
    
    if user_input := st.chat_input("Ask a question about your documents..."):
        handle_query(user_input)

# ------------------------------------------------------------------------------
# 10. Main Entry Point
# ------------------------------------------------------------------------------
def main():
    init_session()
    inject_custom_css()
    
    if not st.session_state.authenticated:
        if st.session_state.get('show_register'):
            register_page()
        else:
            login_page()
    else:
        chat_interface()

if __name__ == "__main__":
    main()